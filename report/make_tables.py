"""Word tables of the exp1 (shuffle control) and exp2 (memory horizon) results.

Run the experiments first so the JSON files exist, then:
    python report/make_tables.py
It writes report/tables/results_tables.docx (A4 landscape):
  1a  exp1 real vs null, log-loss (gain, t, folds won, seed spread, each seed)
  1b  exp1 real vs null, Brier score and AUC
  2   exp2 gain over H=0, lookback EWMAs: log-loss and Brier
  3   exp2 gain over H=0, truncated EWMAs: log-loss and Brier
  4   exp2 log-loss and Brier at each H
All differences are paired per fold over the 10 outer folds; + = better
(lower log-loss / Brier, higher AUC). Needs python-docx.
"""
import json
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "report" / "tables" / "results_tables.docx"
MODEL = {"logreg": "Logistic regression", "gboost": "LightGBM"}
METRIC = {"log_loss": "Log-loss", "brier": "Brier", "auc": "AUC"}
DEC = {"log_loss": 1, "brier": 2, "auc": 1}          # decimals for x1e-4 differences
GROUPS = [(t, m) for t in ("men", "women") for m in ("logreg", "gboost")]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)          # A4 landscape
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.8))
style = doc.styles["Normal"]
style.font.name, style.font.size = "Calibri", Pt(10)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(header_rows, body, widths_cm, first_cols=2):
    t = doc.add_table(rows=0, cols=len(widths_cm))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for hr in header_rows:
        cells = t.add_row().cells
        for c, txt in zip(cells, hr):
            c.text = txt
            shade(c, "D9E2F3")
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True; r.font.size = Pt(9)
    for row in body:
        cells = t.add_row().cells
        for i, (c, txt) in enumerate(zip(cells, row)):
            c.text = txt
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i < first_cols else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(9)
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)
    doc.add_paragraph()
    return t


def caption(title, note, new_page=False):
    p = doc.add_paragraph()
    if new_page:
        p.paragraph_format.page_break_before = True
    r = p.add_run(title); r.bold = True
    if note:
        n = doc.add_paragraph(note); n.runs[0].italic = True; n.runs[0].font.size = Pt(9)


def pm(mean, se, metric):
    """'+x ± y' in units of 1e-4, no '-0.0'."""
    d = DEC[metric]
    mean, se = mean * 1e4, se * 1e4
    if abs(mean) < 0.5 * 10 ** -d:
        mean = 0.0
    return f"{mean:+.{d}f} ± {se:.{d}f}"


def paired(base, other):
    """Mean and s.e. over folds of base - other (+ = other has the lower value)."""
    d = np.asarray(base) - np.asarray(other)
    return d.mean(), d.std(ddof=1) / np.sqrt(len(d))


doc.add_heading("Tennis momentum — results tables", level=1)
doc.add_paragraph(
    "All differences are paired by fold over the 10 temporal outer folds (mean ± standard error), "
    "in units of 10⁻⁴. Positive = better for the first-named condition (lower log-loss or Brier "
    "score, higher AUC). Men and women have separate models.")

# ---- exp1: shuffle control
E1 = json.load(open(REPO / "results" / "exp1_shuffle_control_logreg-gboost.json"))
seeds = [str(s) for s in E1["seeds"]]


def null_mean(tour, m, metric):
    return np.mean([E1["shuffled"][s][tour][m]["summary"][metric]["mean"] for s in seeds])


rows = []
for tour, m in GROUPS:
    P = E1["paired"][tour][m]["log_loss"]
    rows.append([tour.capitalize(), MODEL[m],
                 f"{E1['real'][tour][m]['summary']['log_loss']['mean']:.4f}",
                 f"{null_mean(tour, m, 'log_loss'):.4f}",
                 pm(P["mean"], P["se"], "log_loss"), f"{P['t']:+.1f}", f"{P['better']}/10",
                 f"{P['seed_sd'] * 1e4:.1f}", *[f"{P['per_seed'][s] * 1e4:+.1f}" for s in seeds]])
caption("Table 1a. Shuffle control (exp1): real data vs the order-shuffled null — log-loss.",
        f"Null = points within games and games within sets shuffled, {len(seeds)} seeds "
        f"({', '.join(seeds)}); the null column averages over the seeds. Gain = paired per-fold "
        "improvement of real over the seed-averaged null (+ = point order helps). Seed s.d. = spread of "
        "the gain across seeds; the per-seed columns give each seed's gain (×10⁻⁴). Relational RNN run "
        "in progress.")
table([["Tour", "Model", "Real", "Null", "Gain (×10⁻⁴)", "t", "Real better", "Seed s.d.",
         *[f"Seed {s}" for s in seeds]]],
      rows, [1.6, 3.6, 1.9, 1.9, 2.8, 1.4, 1.8, 1.6, 1.6, 1.6, 1.6])

rows = []
for tour, m in GROUPS:
    row = [tour.capitalize(), MODEL[m]]
    for metric in ("brier", "auc"):
        P = E1["paired"][tour][m][metric]
        row += [f"{E1['real'][tour][m]['summary'][metric]['mean']:.4f}",
                f"{null_mean(tour, m, metric):.4f}",
                pm(P["mean"], P["se"], metric), f"{P['t']:+.1f}", f"{P['better']}/10"]
    rows.append(row)
caption("Table 1b. Shuffle control (exp1): Brier score and AUC.",
        "Same runs as Table 1a. Brier: lower is better, gain = null − real. AUC: higher is better, "
        "gain = real − null. + = point order helps.")
table([["Tour", "Model", "Brier real", "Brier null", "Brier gain (×10⁻⁴)", "t", "Real better",
        "AUC real", "AUC null", "AUC gain (×10⁻⁴)", "t", "Real better"]],
      rows, [1.5, 3.4, 1.8, 1.8, 2.6, 1.2, 1.6, 1.8, 1.8, 2.6, 1.2, 1.6])


# ---- exp2: memory horizon
def horizon(fname, label, num, new_page):
    D = json.load(open(REPO / "results" / f"{fname}.json"))
    hs = list(D["men"]["logreg"]["curve"])
    nmem = [str(D["men"]["logreg"]["curve"][h]["n_memory_features"]) for h in hs[1:]]
    head = ["Tour", "Model", *[("Whole match" if h == "match" else f"H = {h}") for h in hs[1:]]]
    for k, metric in enumerate(("log_loss", "brier")):
        rows = []
        for tour, m in GROUPS:
            c = D[tour][m]["curve"]
            h0 = c["0"]["per_split"][metric]
            rows.append([tour.capitalize(), MODEL[m],
                         *[pm(*paired(h0, c[h]["per_split"][metric]), metric) for h in hs[1:]]])
        caption(f"Table {num}{'ab'[k]}. Memory horizon (exp2, {label}): {METRIC[metric]} gain over the "
                "H = 0 base (×10⁻⁴).",
                "H = 0 is the base model: pre-match and serve features only. Each column adds the "
                "in-match memory features whose lookback is at most H points; 'Whole match' adds all of "
                "them (the full model). + = the extra history helps." if k == 0 else
                f"Same runs as Table {num}a, scored by Brier (lower is better; + = the extra history helps).",
                new_page=new_page and k == 0)
        table([head, ["", "Memory features", *nmem]], rows, [1.6, 3.6] + [2.6] * (len(hs) - 1))
    return D, hs


D2, hs = horizon("exp2_memory_horizon", "lookback EWMAs", 2, new_page=True)
horizon("exp2_memory_horizon_ewma_truncated",
        "truncated EWMAs: every EWMA computed from only the last H points", 3, new_page=True)

rows = []
for metric in ("log_loss", "brier"):
    for tour, m in GROUPS:
        c = D2[tour][m]["curve"]
        rows.append([tour.capitalize(), MODEL[m], METRIC[metric],
                     *[f"{c[h]['summary'][metric]['mean']:.{4 if metric == 'log_loss' else 5}f}" for h in hs]])
caption("Table 4. Memory horizon (exp2, lookback EWMAs): log-loss and Brier score at each history length.",
        "Mean over the 10 folds (lower is better). Unpaired fold-to-fold s.e. is about 0.0010 log-loss "
        "(men) and 0.0005 (women); the paired gains in Tables 2a/2b remove that variation.",
        new_page=True)
table([["Tour", "Model", "Metric", *[("Whole match" if h == "match" else f"H = {h}") for h in hs]]],
      rows, [1.5, 3.3, 1.6] + [2.0] * len(hs), first_cols=3)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("wrote", OUT)
